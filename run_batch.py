# Copyright 2026 Iwamoto-Ai
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
run_batch.py — Excel/CSV の明細（数十件）を、録画リプレイで「1 件ずつ」まとめて処理するバッチランナー。

実務の形（1 回の実行 = 数十件の登録・確認）に合わせた実行層。テンプレートは複雑にせず、
ループ・進捗・エラー処理・再実行は、この実行層が引き受ける。

バッチ定義 JSON の構造:
  {
    "title":  "説明",
    "setup":  [ ...ステップ列... ],   # 最初に 1 回だけ（ログイン〜開始画面まで）
    "loop":   [ ...ステップ列... ],   # 明細 1 行ごとに繰り返す（{{列名}} が行の値で埋まる）
    "recover":[ ...ステップ列... ],   # 失敗した時に開始画面へ戻るための手順（任意）
    "teardown":[ ...ステップ列... ]   # 最後に 1 回（ログアウト等。任意）
  }
  ステップは Chrome Recorder 形式（navigate/click/change/…）に加え、エビデンス保存用の
  {"type":"screenshot","name":"{{project_no}}__{{po_number}}","full_page":false} が使える
  （名前に {{列名}} 可・日時は自動付与 → 「プロジェクト番号__発注番号__日時.png」の命名になる）。

明細ファイル（--details）:
  * CSV（推奨。Excel の「CSV UTF-8」保存）… 1 行目=列名、2 行目以降=データ。列名がそのまま {{キー}}。
  * .xlsx 直読みも可（openpyxl が必要。先頭シート・1 行目=列名。日付は YYYY-MM-DD、整数はそのまま文字列化）。
  * ID 列 … 既定は先頭列。--id-column で変更可。進捗・結果・再実行はこの値で扱う。
  * skip 列 … "skip" 列に何か入っている行は実行せずスキップ（行を消さずに除外できる）。

結果と再実行:
  * 実行のたびに output/batch_result_YYYYMMDD_HHMMSS.csv（ID・結果・理由・エビデンス）と
    output/batch_YYYYMMDD_HHMMSS.log（全ログ）を出力。
  * 1 件の失敗では止まらない（既定）。失敗時はその場のスクショと画面状態を自動保存して次へ。
  * 失敗分だけ再実行: --retry-from output/batch_result_....csv   （または --only ID1,ID2）

使い方（例）:
  python run_batch.py --batch recordings/edi_practice_batch.json --details data/edi_practice_batch.csv --engine playwright --browser edge --no-headless
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import json
import os
import sys

from browser_factory import make_browser
from recorder_import import missing_placeholders
from run_recording import exec_step


# ---------------------------------------------------------------- 明細の読み込み
def _cell_to_str(v) -> str:
    """Excel 由来の値を安全に文字列化する（900000000001.0 のような事故を防ぐ）。"""
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_details(path: str) -> tuple[list[str], list[dict]]:
    """明細ファイル（CSV / xlsx）を読み、(列名リスト, 行 dict のリスト) を返す。"""
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            sys.exit(".xlsx を直接読むには openpyxl が必要です: pip install openpyxl\n"
                     "（または Excel で『CSV UTF-8』として保存し、その CSV を指定してください）")
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [_cell_to_str(h) for h in next(it)]
        rows = []
        for r in it:
            if all(v is None or str(v).strip() == "" for v in r):
                continue
            rows.append({h: _cell_to_str(v) for h, v in zip(headers, r) if h})
        return [h for h in headers if h], rows
    # CSV（utf-8-sig で Excel の BOM を吸収）
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        try:
            headers = [h.strip() for h in next(reader)]
        except StopIteration:
            sys.exit("明細ファイルが空です: " + path)
        rows = []
        for r in reader:
            if not any(c.strip() for c in r):
                continue
            rows.append({h: c.strip() for h, c in zip(headers, r) if h})
    return [h for h in headers if h], rows


def load_failed_ids(result_csv: str, id_col: str) -> list[str]:
    """結果 CSV から失敗した ID の一覧を取り出す（--retry-from 用）。"""
    ids = []
    with open(result_csv, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            if row.get("結果") == "失敗":
                ids.append(row.get(id_col) or row.get("ID") or "")
    return [i for i in ids if i]


# ---------------------------------------------------------------- バッチ実行本体
def run_batch(batch: dict, rows: list[dict], common: dict, browser, out_dir: str,
              id_col: str, log=print, stop_on_error: bool = False,
              full_page: bool = True) -> list[dict]:
    """バッチを実行し、明細ごとの結果（dict のリスト)を返す。
    browser は make_browser() 互換のオブジェクト（テスト時は差し替え可能）。"""
    results: list[dict] = []
    setup = batch.get("setup", [])
    loop = batch.get("loop", [])
    recover = batch.get("recover", [])
    teardown = batch.get("teardown", [])
    total = len(rows)

    # 共通部（ログイン〜開始画面）
    if setup:
        log(f"── セットアップ（{len(setup)} ステップ）──")
        for i, st in enumerate(setup, 1):
            exec_step(browser, st, common, out_dir, full_page=full_page,
                      log=log, tag=f"[setup {i}/{len(setup)}] ")

    aborted = False
    for idx, row in enumerate(rows, 1):
        item_id = row.get(id_col, f"row{idx}")
        vals = {**common, **row}
        rec = {"ID": item_id, "結果": "", "理由": "", "エビデンス": ""}
        # skip 列
        if row.get("skip", "").strip():
            log(f"[{idx}/{total}] {item_id} … スキップ（skip 列指定）")
            rec["結果"] = "スキップ"
            results.append(rec)
            continue
        log(f"── [{idx}/{total}] {item_id} 開始 ──")
        evidence = ""
        try:
            for i, st in enumerate(loop, 1):
                saved = exec_step(browser, st, vals, out_dir, full_page=full_page,
                                  log=log, tag=f"[{idx}/{total} step {i}/{len(loop)}] ")
                if saved:
                    evidence = saved
            rec["結果"] = "成功"
            rec["エビデンス"] = evidence
            log(f"── [{idx}/{total}] {item_id} ✅ 成功 ──")
        except Exception as e:
            rec["結果"] = "失敗"
            rec["理由"] = str(e)[:300]
            rec["エビデンス"] = evidence
            log(f"── [{idx}/{total}] {item_id} ❌ 失敗: {rec['理由']} ──")
            # 失敗時の証跡（その場のスクショと画面状態）を自動保存
            try:
                p = browser.screenshot(os.path.join(out_dir, f"fail_{item_id}.png"),
                                       full_page=full_page)
                log(f"    失敗時スクショ: {p}")
                with open(os.path.splitext(p)[0] + ".txt", "w", encoding="utf-8") as f:
                    f.write(browser.state())
            except Exception:
                pass
            results.append(rec)
            if stop_on_error:
                log("--stop-on-error 指定のため中断します。")
                aborted = True
                break
            # 開始画面へ戻る（recover）。ここで失敗したら以降は続行不能とみなし中断。
            if recover:
                try:
                    log("    復帰手順（recover）を実行 …")
                    for st in recover:
                        exec_step(browser, st, vals, out_dir, full_page=full_page,
                                  log=log, tag="    [recover] ")
                except Exception as re:
                    log(f"復帰にも失敗したため中断します: {str(re)[:200]}")
                    aborted = True
                    break
            continue
        results.append(rec)

    if teardown and not aborted:
        log(f"── 後片付け（{len(teardown)} ステップ）──")
        for st in teardown:
            try:
                exec_step(browser, st, common, out_dir, full_page=full_page,
                          log=log, tag="[teardown] ")
            except Exception:
                pass
    return results


# ---------------------------------------------------------------- CLI
def main() -> None:
    p = argparse.ArgumentParser(
        description="Excel/CSV の明細を録画リプレイでまとめて処理するバッチランナー（LLM 不要）")
    p.add_argument("--batch", required=True, help="バッチ定義 JSON（setup/loop/recover/teardown）")
    p.add_argument("--details", required=True, help="明細ファイル（CSV 推奨 / .xlsx 可）")
    p.add_argument("--values", default=None, help="全件共通の値 JSON（任意。行の値が優先）")
    p.add_argument("--id-column", default=None, help="ID に使う列名（既定: 先頭列）")
    p.add_argument("--only", default=None, help="この ID だけ実行（カンマ区切り）")
    p.add_argument("--retry-from", default=None, help="結果 CSV を指定し、失敗分だけ再実行")
    p.add_argument("--stop-on-error", action="store_true", help="最初の失敗で中断する（既定: 続行）")
    p.add_argument("--max-items", type=int, default=0, help="先頭から N 件だけ実行（0=全件。お試し用）")
    p.add_argument("--engine", choices=["selenium", "playwright"], default="playwright")
    p.add_argument("--browser", choices=["edge", "chrome"], default="edge")
    p.add_argument("--out-dir", default="output")
    p.add_argument("--viewport-shot", dest="full_page", action="store_false",
                   help="スクショを表示領域（横長）だけにする。既定はページ全体")
    p.set_defaults(full_page=True)
    g = p.add_mutually_exclusive_group()
    g.add_argument("--headless", dest="headless", action="store_true")
    g.add_argument("--no-headless", dest="headless", action="store_false")
    p.set_defaults(headless=False)
    args = p.parse_args()

    with open(args.batch, encoding="utf-8") as f:
        batch = json.load(f)
    common = {}
    if args.values:
        with open(args.values, encoding="utf-8") as f:
            common = json.load(f)

    headers, rows = load_details(args.details)
    if not rows:
        sys.exit("明細が 0 件です: " + args.details)
    id_col = args.id_column or headers[0]
    if id_col not in headers:
        sys.exit(f"ID 列 '{id_col}' が明細の列名にありません: {headers}")

    # 対象の絞り込み（--only / --retry-from）
    if args.retry_from:
        only = set(load_failed_ids(args.retry_from, id_col))
        if not only:
            sys.exit("再実行対象（失敗）が結果 CSV にありません: " + args.retry_from)
        rows = [r for r in rows if r.get(id_col) in only]
        print(f"再実行対象: {len(rows)} 件（{args.retry_from} の失敗分）")
    elif args.only:
        only = {s.strip() for s in args.only.split(",") if s.strip()}
        rows = [r for r in rows if r.get(id_col) in only]
        print(f"対象を絞り込み: {len(rows)} 件（--only）")
    if args.max_items and args.max_items > 0:
        rows = rows[: args.max_items]
    if not rows:
        sys.exit("実行対象が 0 件です（--only / --retry-from の指定を確認してください）")

    # プレースホルダ検査（明細 1 行目＋共通値で loop を検査）
    sample = {**common, **rows[0]}
    miss = missing_placeholders({"steps": batch.get("loop", [])}, sample)
    miss += missing_placeholders({"steps": batch.get("setup", [])}, common)
    if miss:
        sys.exit("値が未指定のプレースホルダがあります: " + ", ".join(sorted(set(miss)))
                 + "\n明細の列名（CSV 1 行目）か --values の JSON に追加してください。")

    os.makedirs(args.out_dir, exist_ok=True)
    ts = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = os.path.join(args.out_dir, f"batch_{ts}.log")
    result_path = os.path.join(args.out_dir, f"batch_result_{ts}.csv")
    logf = open(log_path, "w", encoding="utf-8")

    def log(msg: str) -> None:
        print(msg)
        logf.write(msg + "\n")
        logf.flush()

    log(f"バッチ: {batch.get('title', args.batch)}")
    log(f"明細: {args.details}（{len(rows)} 件 / ID列: {id_col}）")

    browser = make_browser(args.engine, args.browser, args.headless)
    results = None
    try:
        results = run_batch(batch, rows, common, browser, args.out_dir, id_col,
                            log=log, stop_on_error=args.stop_on_error,
                            full_page=args.full_page)
    except Exception as e:
        log(f"\n❌ 続行できないエラーで中断しました: {str(e)[:300]}")
        log("   よくある原因: (1) {{SECRET:...}} の環境変数が未設定（実行前に $env:名前=\"値\" を設定）"
            "  (2) 練習サイトが未配信  (3) 開始 URL に到達できない")
    finally:
        browser.quit()
    if results is None:
        logf.close()
        sys.exit(2)

    # 結果 CSV（Excel で開けるよう BOM 付き UTF-8）
    with open(result_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([id_col, "結果", "理由", "エビデンス"])
        for r in results:
            w.writerow([r["ID"], r["結果"], r["理由"], r["エビデンス"]])

    ok = sum(1 for r in results if r["結果"] == "成功")
    ng = sum(1 for r in results if r["結果"] == "失敗")
    sk = sum(1 for r in results if r["結果"] == "スキップ")
    log(f"\n===== 結果: {len(results)} 件中  成功 {ok} / 失敗 {ng} / スキップ {sk} =====")
    if ng:
        log("失敗した明細: " + ", ".join(r["ID"] for r in results if r["結果"] == "失敗"))
        log(f"失敗分だけ再実行するには: python run_batch.py --batch {args.batch} "
            f"--details {args.details} --retry-from {result_path}")
    log(f"結果 CSV: {result_path}")
    log(f"ログ: {log_path}")
    logf.close()
    sys.exit(1 if ng else 0)


if __name__ == "__main__":
    main()
